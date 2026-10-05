"""
Created on Mon Apr  5 09:54:58 2021
@author: ZhengBin
"""
import sys
sys.path.append('../')
import os
import copy
import time
import numpy as np
import pickle
import xlwt
import pandas as pd
from src.jpkfile import JPKFile, JPKMap
import zipfile
from zipfile import ZipFile
from scipy.signal import savgol_filter
import pathlib
import lzma
from datayee.resources import asset_path
'''from nanoscope import files
from nanoscope.constants import FORCE, METRIC, VOLTS, PLT_kwargs'''
def rotate(data_x, data_y, index, k):
    theta = np.arctan(k) * -1
    if type(index) == int:
        return (data_x - data_x[index]) * np.sin(theta) + (data_y - data_y[index]) * np.cos(theta) + data_y[index]
    elif type(index) == tuple:
        return (data_x - index[0]) * np.sin(theta) + (data_y - index[1]) * np.cos(theta) + index[1]
    
    


class forcecurve:
    def __init__(self):
        self.data = {'tasktype': '',
                     'rawdata': {},
                     'path': '',
                     'springConstant': 0.01,
                     'datamsg': ('', 0),
                     'offset': {'x': 0, 'y': 0, 'k': 0,'highspeed':0,'k_alpha':1},
                     'filters': {'methods': 'savgol', 'win_lens': 13, 'poly': 2},
                     'mobilenet_judge': True,
                     'peaknum_judge': True,
                     'artificial_judge': True,
                     'peakindex': [],
                     'bottomindex': [],
                     'wlcarg': [],
                     'dlc': [],
                     'k': [],
                     'mark': [],
                     'arg':{},
                     'xy-position':[0,0],
                     'compressed_data':{},
                     'class':'N',
                     'springConstant_override': None,
                     'overlay':False,
                     'wlc_err': [],
                     'peak_cluster': [],
                     'cluster_pca': None,
                     'peak_forces': []}
        self._prodata_cache = {}
        self._prodata_dirty = True

    def invalidate_prodata_cache(self):
        self._prodata_cache.clear()
        self._prodata_dirty = True

    def get_prodata(self, smooth=True, tip_correc=True, s=None):
        cache_key = (smooth, tip_correc, -1 if s is None else s)
        if not self._prodata_dirty and cache_key in self._prodata_cache:
            return self._prodata_cache[cache_key]
        data = copy.deepcopy(self.data['rawdata'])
        if 'k_alpha' not in self.data['offset'].keys():
            self.data['offset']['k_alpha'] = 1
        for k, v in data.items():
            data[k]['vDeflection'] *= self.data['offset']['k_alpha']
            ox = self.data['offset']['x']
            if isinstance(ox, np.ndarray):
                ox = ox.item() if ox.ndim == 0 else ox.reshape(-1)[0]
            data[k]['measuredHeight'] = data[k]['measuredHeight'] - float(ox)
            data[k]['vDeflection'] = data[k]['vDeflection'] - self.data['offset']['y']
            if smooth and k == 'retract':
                if s == None:
                    data[k]['vDeflection'] = savgol_filter(data[k]['vDeflection'][:, 0],
                                                           self.data['filters']['win_lens'],
                                                           self.data['filters']['poly']).reshape(
                        len(data[k]['vDeflection']), 1)
                else:
                    data[k]['vDeflection'] = savgol_filter(data[k]['vDeflection'][:, 0], s, 2).reshape(
                        len(data[k]['vDeflection']), 1)
            data[k]['vDeflection'] *= -1
            if tip_correc:
                sc = self.data['springConstant_override'] if self.data.get('springConstant_override') is not None else self.data['springConstant']
                data[k]['measuredHeight'] = data[k]['measuredHeight'] - (data[k]['vDeflection']-data[k]['vDeflection'].min()) / sc
            if 'k' in self.data['offset'].keys():
                if 'rotate_index' in self.data['offset'].keys():
                    rotate_index = self.data['offset']['rotate_index']
                else:
                    rotate_index = -1
                rotate_x,rotate_y = data['retract']['measuredHeight'][rotate_index],data['retract']['vDeflection'][rotate_index]
                data[k]['vDeflection'] = rotate(data[k]['measuredHeight'].reshape(-1),
                                                data[k]['vDeflection'].reshape(-1),
                                                (rotate_x,rotate_y),
                                                self.data['offset']['k']).reshape(-1, 1)
        self._prodata_cache[cache_key] = data
        self._prodata_dirty = False
        return data

    def savedata2txt(self, savedir='data.txt'):
        data = self.get_prodata(tip_correc=False)
        data_re = data['retract']
        f = data_re['vDeflection'].reshape(-1)
        h = data_re['measuredHeight'].reshape(-1)
        data2save = np.dstack((f, h))[0]
        now = int(time.time())
        timeArray = time.localtime(now)
        otherStyleTime = time.strftime("%Y-%m-%d %H:%M:%S", timeArray)
        header = ''
        header += otherStyleTime
        header += '\n'
        header += 'springConstant:' + str(self.data['springConstant']) + 'N/m'
        with open(savedir, 'w') as f:
            np.savetxt(f, data2save, header=header, fmt='%.6e')
    def compress(self):
        self.data['compressed_data']={}
        for item,value in self.data['rawdata'].items():
            self.data['compressed_data'][item] = (lzma.compress(value.astype([('measuredHeight', '<f4'), ('vDeflection', '<f4')]).tobytes()),eval(str(value.dtype).replace("<f8","<f4")))
    def decompress(self):
        for item,value in self.data['compressed_data'].items():
            self.data['rawdata'][item] = np.frombuffer(lzma.decompress(value[0]),value[1]).reshape(-1,1)
    def clean_force(self):
        self.data['rawdata'] = {}
        self._prodata_cache.clear()
        self._prodata_dirty = True
    def recover_force(self, ljf=None):
        if len(self.data['rawdata'])!=0:
            return None
        self.data['rawdata']= {}
        if 'compressed_data' not in self.data.keys() or self.data['compressed_data']=={}:
            if ljf==None:
                return False
            ljf.file_type_deter(*self.data['datamsg'])
            self.data['rawdata'] = ljf.data['rawdata']
            self.compress()
        self.decompress()
        self._prodata_cache.clear()
        self._prodata_dirty = True
        return True
class loadjpkfile():
    def __init__(self,filedir):
        self.filedir = filedir
        self.filelst = []
        self.datalst = []
        self.get_filelst()
        self.get_datalst()
        self.startnum = -1
        self.current_readfile = 'none'
        fc = forcecurve()
        self.data = fc.data
        self.data['path'] = self.filedir
        self.buffer = {'txt':'none','jpkmap':'none','jpkforce':'nono','datay':'none'}
        self.datatype = [('measuredHeight', '<f4'), ('vDeflection', '<f4')]
    def __len__(self):
        return len(self.datalst)
    def __iter__(self):
        return self
    def __next__(self):
        self.startnum += 1
        if self.startnum < len(self):
            pass
        else:
            self.startnum = -1
            raise StopIteration
        self.data['datamsg'] = self.datalst[self.startnum]
        self.file_type_deter(*self.datalst[self.startnum])
        return copy.deepcopy(self.data)
    def __getitem__(self, index):
        self.file_type_deter(*self.datalst[index])
        self.data['datamsg'] = self.datalst[index]
        return copy.deepcopy(self.data)
    def get_datalst(self):
        for fname in self.filelst:
            if sum([True for i in ['.txt', '.jpk-force', '.datay'] if fname.endswith(i)]):
                self.datalst.append((fname, 0))
            elif sum([True for i in ['.jpk-force-map'] if fname.endswith(i)]):
                properties = ZipFile(fname).open('header.properties')
                while True:
                    line = properties.readline()
                    if b'force-scan-map.indexes.max' in line:
                        maxindex = int(line.rstrip().split(b'=')[-1])
                        for i in range(maxindex):
                            self.datalst.append((fname, i))
                        break
                        '''
            elif sum([True for i in ['.spm'] if fname.endswith(i)]):
                with files.ForceVolumeFile(fname) as f:
                    fv_pixels = f.force_curves_channel.number_of_force_curves
                    for i in range(fv_pixels):
                        self.datalst.append((fname, i))
                        '''
    def get_filelst(self, Travel=True):
        if os.path.isfile(self.filedir):
            self.filelst.append(self.filedir)
        elif os.path.isdir(self.filedir):
            for a, b, c in os.walk(self.filedir, topdown=True, onerror=None, followlinks=False):
                for filename in c:
                    if sum([True for i in ['.txt', '.jpk-force', '.jpk-force-map', '.datay','.spm'] if
                            os.path.join(a, filename).endswith(i)]):
                        self.filelst.append(os.path.join(a, filename))
                if not Travel:
                    break
    def file_type_deter(self, filename, index):
        self.change_buffer(filename)
        if filename.endswith('.txt'):
            self.extract_txt_data(filename, index)
        elif filename.endswith('.jpk-force'):
            self.extract_force_data(filename, index)
        elif filename.endswith('.jpk-force-map'):
            self.extract_map_data(filename, index)
        elif filename.endswith('.datay'):
            self.extract_datay_data(filename, index)
        elif filename.endswith('.spm'):
            self.extract_spm_data(filename, index)
    def change_buffer(self,filename):
        if self.current_readfile != filename:
            self.current_readfile = filename
            suffix = os.path.splitext(filename)[-1]
            if suffix =='.txt':
                self.buffer['txt'] = np.loadtxt(filename, comments='#')
            elif suffix =='.jpk-force':
                self.buffer['jpkforce'] = JPKFile(filename)
            elif suffix == '.jpk-force-map':
                self.buffer['jpkmap'] = JPKMap(filename)
            elif suffix == '.datay':
                with open(filename,'rb') as f:
                    self.buffer['datay'] = pickle.load(f)
            '''
            elif suffix == '.spm':
                self.buffer['spm'] = []
                with files.ForceVolumeFile(filename) as f:
                    fc_channel = f.force_curves_channel
                    fv_pixels = f.force_curves_channel.number_of_force_curves
                    h_sens_chan = f[2]
                    for index in range(fv_pixels):
                        fz_plot, ax_prop = fc_channel.create_force_z_plot(index, FORCE)
                        h_sens_data = h_sens_chan.get_force_curve_data(index, METRIC)
                        d_ = copy.deepcopy(self.data)
                        if 'nN' in ax_prop['ylabel']:
                            factor = 1e-9
                        elif 'pN' in ax_prop['ylabel']:
                            factor = 1e-12
                        for c in ['retract','extend']:
                            if c =='retract':
                                data_x = h_sens_data.retrace * -1e-9
                                data_y = fz_plot.retrace.y * factor
                            elif c == 'extend':
                                data_x = h_sens_data.trace * -1e-9
                                data_y = fz_plot.trace.y * factor
                            data = np.dstack((data_x, data_y))[0]
                            d_['rawdata'][c] = np.array([[tuple(i)] for i in data],
                                                            dtype=[('measuredHeight', '<f4'), ('vDeflection', '<f4')])
                        d_['springConstant'] = f.spring_constant
                        self.buffer['spm'].append(d_)
            '''
    def extract_txt_data(self, filename, index):
        data = self.buffer['txt']
        with open(filename, 'r') as f:
            text = f.readlines()
            for line in text:
                if '# springConstant' in line:
                    springConstant = float(line.split()[-1])
                    break
        if 'springConstant' not in locals().keys():
            if '# springConstant' in line:
                springConstant = float(line.replace('"','').split()[-1])
            else:
                springConstant = 0.01
        self.data['springConstant'] = springConstant
        self.data['rawdata']['extend'] = np.array([[tuple(i)] for i in data[:np.argmin(data[:, 0])]],
                                                  dtype=self.datatype)
        self.data['rawdata']['retract'] = np.array([[tuple(i)] for i in data[np.argmin(data[:, 0]):]],
                                                   dtype=self.datatype)
    def extract_spm_data(self, fname, index):
        self.data = self.buffer['spm'][index]
        '''
        with files.ForceVolumeFile(fname) as f:
            print(fname)
            fc_channel = f.force_curves_channel
            h_sens_chan = f[2]
            fz_plot, ax_prop = fc_channel.create_force_z_plot(index, FORCE)
            h_sens_data = h_sens_chan.get_force_curve_data(index, METRIC)
            if 'nN' in ax_prop['ylabel']:
                factor = 1e-9
            elif 'pN' in ax_prop['ylabel']:
                factor = 1e-12
            data_x = h_sens_data.retrace * -1e-9
            data_y = fz_plot.retrace.y * factor
            data = np.dstack((data_x, data_y))[0]
            self.data['rawdata']['retract'] = np.array([[tuple(i)] for i in data],
                                                       dtype=[('measuredHeight', '<f4'), ('vDeflection', '<f4')])
            self.data['springConstant'] = f.spring_constant
        '''
    def extract_force_data(self, filename, index):
        jpk = self.buffer['jpkforce']
        springConstant = 0
        for i in ['1','2']:
            try:
                springConstant = float(
                jpk.shared_parameters['lcd-info'][i]['conversion-set']['conversion']['force']['scaling'][
                    'multiplier'])
                break
            except:
                continue
        if springConstant == 0:
            return None
        self.data['springConstant'] = springConstant
        for i, segment in jpk.segments.items():
            self.data['rawdata'][segment.get_info('type')] = segment.get_array(['measuredHeight', 'vDeflection'])[0]
    def extract_map_data(self, filename, index):
        jpks = self.buffer['jpkmap']
        jpk = jpks.get_single_pixel(index)
        position = jpks.flat_indices[index].parameters['force-scan-series']['header']['position']
        self.data['xy-position'][0],self.data['xy-position'][1] = float(position['x']),float(position['y'])
        try:
            springConstant = float(
                jpk.shared_parameters['lcd-info']['2']['conversion-set']['conversion']['force']['scaling'][
                    'multiplier'])
        except:
            springConstant = float(
                jpk.shared_parameters['lcd-info']['1']['conversion-set']['conversion']['force']['scaling'][
                    'multiplier'])
        self.data['springConstant'] = springConstant
        for i, segment in jpk.segments.items():
            self.data['rawdata'][segment.get_info('type')] = segment.get_array(['measuredHeight', 'vDeflection'])[0]
    def extract_datay_data(self, filename, index):
        self.data['springConstant'] = self.buffer['datay']['springConstant']
        self.data['rawdata'] = self.buffer['datay']['rawdata']
class loadjpkfile_(forcecurve):
    def __init__(self, filedir):
        super().__init__()
        self.filedir = filedir
        self.filelst = []
        self.datalst = []
        self.get_filenamelst()
        self.get_dataindex()
        self.startnum = -1
        self.data_structure = copy.deepcopy(self.data)

    def __len__(self):
        return len(self.datalst)

    def __iter__(self):
        return self

    def __next__(self):
        self.startnum += 1
        if self.startnum < len(self.datalst):
            pass
        else:
            self.startnum = -1
            raise StopIteration
        self.data = copy.deepcopy(self.data_structure)
        self.file_type_deter(*self.datalst[self.startnum])
        self.data['datamsg'] = self.datalst[self.startnum]
        self.data['path'] = self.filedir
        return self.data

    def __getitem__(self, index):
        self.data = copy.deepcopy(self.data_structure)
        self.file_type_deter(*self.datalst[index])
        self.data['datamsg'] = self.datalst[index]
        self.data['path'] = self.filedir
        return self.data

    def get_dataindex(self):
        for fname in self.filelst:
            if sum([True for i in ['.txt', '.jpk-force', '.datay'] if fname.endswith(i)]):
                self.datalst.append((fname, 0))
            elif sum([True for i in ['.jpk-force-map'] if fname.endswith(i)]):
                properties = ZipFile(fname).open('header.properties')
                while True:
                    line = properties.readline()
                    if b'force-scan-map.indexes.max' in line:
                        maxindex = int(line.rstrip().split(b'=')[-1])
                        for i in range(maxindex):
                            self.datalst.append((fname, i))
                        break
            '''
            elif sum([True for i in ['.spm'] if fname.endswith(i)]):
                with files.ForceVolumeFile(fname) as f:
                    fv_pixels = f.force_curves_channel.number_of_force_curves
                    for i in range(fv_pixels):
                        self.datalst.append((fname, i))
            '''
    def get_filenamelst(self, Travel=True):
        if os.path.isfile(self.filedir):
            self.filelst.append(self.filedir)
        elif os.path.isdir(self.filedir):
            for a, b, c in os.walk(self.filedir, topdown=True, onerror=None, followlinks=False):
                for filename in c:
                    if sum([True for i in ['.txt', '.jpk-force', '.jpk-force-map', '.datay'] if
                            os.path.join(a, filename).endswith(i)]):
                        self.filelst.append(os.path.join(a, filename))
                if not Travel:
                    break

    def file_type_deter(self, filename, index):
        if filename.endswith('.txt'):
            self.extract_txt_data(filename, index)
        elif filename.endswith('.jpk-force'):
            self.extract_force_data(filename, index)
        elif filename.endswith('.jpk-force-map'):
            self.extract_map_data(filename, index)
        elif filename.endswith('.datay'):
            self.extract_datay_data(filename, index)
        '''
        elif filename.endswith('.spm'):
            self.extract_spm_data(filename, index)'''

    def extract_txt_data(self, filename, index):
        data = np.loadtxt(filename, comments='#')
        with open(filename, 'r') as f:
            text = f.readlines()
            for line in text:
                if '# springConstant' in line:
                    springConstant = float(line.split()[-1])
                    break
        self.data['springConstant'] = springConstant
        self.data['rawdata']['extend'] = np.array([[tuple(i)] for i in data[:np.argmin(data[:, 0])]],
                                                  dtype=self.datatype)
        self.data['rawdata']['retract'] = np.array([[tuple(i)] for i in data[np.argmin(data[:, 0]):]],
                                                   dtype=self.datatype)

    def extract_force_data(self, filename, index):
        try:
            jpk = JPKFile(filename)
        except:
            return None
        try:
            springConstant = float(
                jpk.shared_parameters['lcd-info']['2']['conversion-set']['conversion']['force']['scaling'][
                    'multiplier'])
        except:
            return None
        self.data['springConstant'] = springConstant
        for i, segment in jpk.segments.items():
            self.data['rawdata'][segment.get_info('type')] = segment.get_array(['measuredHeight', 'vDeflection'])[0]

    def extract_map_data(self, filename, index):
        try:
            jpks = JPKMap(filename)
        except:
            return None
        jpk = jpks.get_single_pixel(index)
        position = jpks.flat_indices[index].parameters['force-scan-series']['header']['position']
        self.data['xy-position'][0],self.data['xy-position'][1] = float(position['x']),float(position['y'])
        try:
            springConstant = float(
                jpk.shared_parameters['lcd-info']['2']['conversion-set']['conversion']['force']['scaling'][
                    'multiplier'])
        except:
            springConstant = float(
                jpk.shared_parameters['lcd-info']['1']['conversion-set']['conversion']['force']['scaling'][
                    'multiplier'])
        self.data['springConstant'] = springConstant
        for i, segment in jpk.segments.items():
            self.data['rawdata'][segment.get_info('type')] = segment.get_array(['measuredHeight', 'vDeflection'])[0]
    '''
    def extract_spm_data(self, fname, index):
        with files.ForceVolumeFile(fname) as f:
            fc_channel = f.force_curves_channel
            h_sens_chan = f[2]
            fz_plot, ax_prop = fc_channel.create_force_z_plot(index, FORCE)
            h_sens_data = h_sens_chan.get_force_curve_data(index, METRIC)
            if 'nN' in ax_prop['ylabel']:
                factor = 1e-9
            elif 'pN' in ax_prop['ylabel']:
                factor = 1e-12
            data_x = h_sens_data.retrace * -1e-9
            data_y = fz_plot.retrace.y * factor
            data = np.dstack((data_x, data_y))[0]
            self.data['rawdata']['retract'] = np.array([[tuple(i)] for i in data],
                                                       dtype=[('measuredHeight', '<f4'), ('vDeflection', '<f4')])
            self.data['springConstant'] = f.spring_constant
    '''
    def extract_all_map2datay(self, todir):
        dic = self.data_structure
        for filename in self.filelst:
            if filename.endswith('.jpk-force-map'):
                jpks = JPKMap(filename)
                for i,j in jpks.flat_indices.items():

                    jpk = jpks.get_single_pixel(i)
                    position = j.parameters['force-scan-series']['header']['position']
                    dic['xy-position'][0],self.data['xy-position'][1] = float(position['x']),float(position['y'])
                    try:
                        springConstant = float(
                            jpk.shared_parameters['lcd-info']['2']['conversion-set']['conversion']['force']['scaling'][
                                'multiplier'])
                    except:
                        springConstant = float(
                            jpk.shared_parameters['lcd-info']['1']['conversion-set']['conversion']['force']['scaling'][
                                'multiplier'])
                    dic['springConstant'] = springConstant
                    dic['datamsg'] = (filename, i)
                    for n, segment in jpk.segments.items():
                        dic['rawdata'][segment.get_info('type')] = segment.get_array(['measuredHeight', 'vDeflection'])[
                            0]
                    fname = "{}-{}.datay".format(os.path.join(todir, os.path.splitext(os.path.basename(filename))[0]),
                                                 i)
                    with open(fname, 'wb') as f:
                        pickle.dump(dic, f)
            elif filename.endswith('.jpk-force'):
                jpk = JPKFile(filename)
                try:
                    springConstant = float(
                        jpk.shared_parameters['lcd-info']['2']['conversion-set']['conversion']['force']['scaling'][
                            'multiplier'])
                except:
                    continue
                dic['springConstant'] = springConstant
                fname = "{}-{}.datay".format(os.path.join(todir, os.path.splitext(os.path.basename(filename))[0]),0)
                for i, segment in jpk.segments.items():
                    dic['rawdata'][segment.get_info('type')] = segment.get_array(['measuredHeight', 'vDeflection'])[0]
                with open(fname, 'wb') as f:
                    pickle.dump(dic, f)

    def extract_datay_data(self, filename, index):
        with open(filename, 'rb') as f:
            pkl = pickle.load(f)
        springConstant = pkl['springConstant']
        self.data['springConstant'] = springConstant
        self.data['rawdata'] = pkl['rawdata']


class zipfileopera:
    def __init__(self, fname='test.DataYee-force'):
        self.fname = fname
        self.startnum = -1
        self.version = 'version2'
        self.data = {self.version:'','data.pkl':{}}
        self.readfile()
        self.change = {}
        self.delete = []

    def __len__(self):
        return len(self.data['data.pkl'])

    def __getitem__(self, index):
        index = list(self.data['data.pkl'].keys())[index]
        return self.data['data.pkl'][index]
    def readfile(self):
        if os.path.isfile(self.fname):
            with ZipFile(self.fname, 'r', zipfile.ZIP_DEFLATED) as zips:
                if self.version in zips.namelist():
                    for fname in zips.namelist():
                        with zips.open(fname) as f:
                            self.data[fname] = pickle.load(f)
                else:
                    for fname in zips.namelist():
                        if os.path.splitext(fname)[-1] == '.pkl':
                            with zips.open(fname) as f:
                                fcdata = pickle.load(f)
                                self.data['data.pkl'][fcdata['datamsg']] = fcdata
    def savefile(self):
        with ZipFile(self.fname, 'w', zipfile.ZIP_DEFLATED) as zips:
            for k,v in self.data.items():
                zips.writestr(k, pickle.dumps(v))
    def get_sourcepath(self):
        return list(self.data['data.pkl'].values())[0]['path']
    def addforce(self, fc):
        if len(fc.data.get('rawdata', {})) > 0:
            fc.compress()
        fc.clean_force()
        self.data['data.pkl'][fc.data['datamsg']] = fc.data
    def changingforce(self, fc):
        if len(fc.data.get('rawdata', {})) > 0:
            fc.compress()
        fc.clean_force()
        self.change[fc.data['datamsg']] = copy.deepcopy(fc.data)
    def deletingforce(self,fc):
        self.delete.append(fc.data['datamsg'])
    def deletedforce(self):
        for d in self.delete:
            if d in self.data['data.pkl'].keys():
                del self.data['data.pkl'][d]
        self.delete = []
        self.savefile()
    def clean_force(self):
        self.data['data.pkl'] = {}
    def changedforce(self, svfname='123.DataYee-force',saveas=False,save=True):
        if saveas:
            self.fname = svfname
        for k,v in self.change.items():
            self.data['data.pkl'][k] = v
        self.change = {}
        if save:
            self.savefile()
    def delet_dataYee(self):
        self.data = {self.version:'','data.pkl':{}}
        if os.path.isfile(self.fname):
            os.remove(self.fname)
    def get_maxforce(self,ljp,filters=True, filter_lst=['peaknum_judge', 'mobilenet_judge', 'artificial_judge']):
        arr = np.array([])
        for i in range(len(self)):
            fc1 = forcecurve()
            fc1.data = self[i]
            fc1.recover_force(ljp)
            data = fc1.get_prodata()['retract']
            data_y = data['vDeflection']*1e12
            if len(fc1.data['peakindex'])==0:
                max_d = data_y[int(0.8*len(data_y)):].max()
            else:
                max_d = data_y[fc1.data['peakindex']].max()
            if max_d<0:
                max_d = 0
            arr = np.append(arr,max_d)
        with open('maxforce.txt','w') as f:
            np.savetxt(f,arr)
        return arr
    def get_prodata(self,classes=[]):
        dic = {}
        dic['Rupforce'] = {}
        for data in self:
            if data['artificial_judge'] and data['class'] not in ['N','F']:
                fc  =forcecurve()
                fc.data = data
                mark  = fc.data['mark']
                p = fc.data['peakindex']
                if len(classes)!=0 and fc.data['class'] not in classes:
                    continue
                if fc.recover_force():
                    pass
                else:
                    ljp = loadjpkfile(self.get_sourcepath())
                    fc.recover_force(ljp)
                d = fc.get_prodata()['retract']
                x,y = d['measuredHeight']*1e9,d['vDeflection']*1e12
                for i,m in enumerate(mark):
                    if m not in dic.keys():
                        dic[m] = dict(force=[],dlc=[],lc=[],k=[])
                    dic[m]['force'].append((y[p[i]])[0])
                    dic[m]['dlc'].append(fc.data['dlc'][i])
                    dic[m]['lc'].append(fc.data['wlcarg'][i][0])
                    dic[m]['k'].append(fc.data['k'][i])
                if fc.data['class'] not in dic['Rupforce'].keys():
                    dic['Rupforce'][fc.data['class']] = dict(force=[],lc=[],k=[])
                dic['Rupforce'][fc.data['class']]['force'].append(y[p[-1]][0])
                dic['Rupforce'][fc.data['class']]['lc'].append(fc.data['wlcarg'][-1][0])
                dic['Rupforce'][fc.data['class']]['k'].append(fc.data['k'][-1])
                fc.clean_force()
        return dic
                    
    def resortForce(self,index_lst):
        if len(index_lst)<max(index_lst) or len(self)!=len(index_lst):
            return False
        keys = list(self.data['data.pkl'].keys())
        sorted_keys = [keys[i] for i in index_lst]
        new_pkl = {}
        for k in sorted_keys:
            new_pkl[k] = self.data['data.pkl'][k]
        self.data['data.pkl'] = new_pkl
        self.savefile()
    def split_bypeakN(self):
        fc = forcecurve()
        dic = {}
        pure_fname = os.path.splitext(self.fname)[0]
        for i,data in enumerate(self):
            fc.data = data
            peakN = fc.data['peakindex']
            if len(peakN) not in dic.keys():
                dic[peakN] = []
            dic[peakN].append(fc)
        for k,v in dic.items():
            fname = '{}-peakN-{}.DataYee-Force'.format(pure_fname,k)
            with ZipFile(fname, 'a', zipfile.ZIP_DEFLATED) as zips:
                for fc in v:
                    o = os.path.splitext(os.path.basename(fc.data['datamsg'][0]))[0] + '-s-' + str(
                        fc.data['datamsg'][1]) + '.pkl'
                    pkl = pickle.dumps(fc.data)
                    zips.writestr(o, pkl)
    def split_DataYee(self,SplitDic):
        rawname = os.path.splitext(os.path.basename(self.fname))[0]
        dirname = os.path.join(os.path.dirname(self.fname),'clustering')
        if not os.path.isdir(dirname):
            os.makedirs(dirname)
        with ZipFile(self.fname, 'r', zipfile.ZIP_DEFLATED) as zips:
            lst = copy.deepcopy(zips.namelist())
        with ZipFile(self.fname, 'r', zipfile.ZIP_DEFLATED) as zips:
            for classindex,indexlst in SplitDic.items():
                outputname = os.path.join(dirname, '{}-{}.{}'.format(rawname,classindex,'DataYee-force'))
                if os.path.isfile(outputname):
                    os.remove(outputname)
                for index in indexlst:
                    arcname = lst[index]
                    with zips.open(arcname, 'r') as f:
                        pkl = f.read()
                    with ZipFile(outputname, 'a', zipfile.ZIP_DEFLATED) as zips1:
                        zips1.writestr(arcname, pkl)                
        
    def exporttxt(self,ljp,forcecurve_index,tip_correc=True):
        fc = forcecurve()
        fc.data=self[forcecurve_index]
        if not fc.data['artificial_judge']:
            return None
        fc.recover_force(ljp)
        data = fc.get_prodata(tip_correc=tip_correc)
        data_x = data['retract']['measuredHeight']
        data_y = data['retract']['vDeflection']*-1
        x = np.dstack((data_x[:,0],data_y[:,0]))[0]
        if 'extend' in data.keys():
            data_x_e = data['extend']['measuredHeight']
            data_y_e = data['extend']['vDeflection']*-1
            e = np.dstack((data_x_e[:,0],data_y_e[:,0]))[0]
            x = pd.DataFrame(np.vstack((e,np.array([np.nan,np.nan]),x)))
        else:
            x = pd.DataFrame(np.vstack((np.array([[0,0],[np.nan,np.nan]]),x)))
        todir = os.path.join(os.path.dirname(self.fname),'txt_out')
        if not os.path.isdir(todir):
            os.makedirs(todir)
        name = "{}.txt".format(forcecurve_index)
        fname = os.path.join(todir, name)
        header = ['#','SpringConstant: {:.4f}'.format(fc.data['springConstant'])]
        x.to_csv(fname,sep=' ',float_format='%.5e',index=False,header=header)
    def get_arg(self,ljp,f_index=None,include_last=True):
        arg_dic = {'dlc':[],'lc':[],'p':[],'force':[],'k':[],'lens':[]}
        mark_data = {}

        def append_mark(mark, curve_index, peak_index, data_index, dlc, lc, p, k, force):
            if mark not in mark_data:
                mark_data[mark] = {
                    'curve_index': [],
                    'peak_number': [],
                    'data_index': [],
                    'dlc': [],
                    'lc': [],
                    'p': [],
                    'k': [],
                    'force': []
                }
            values = (curve_index, peak_index + 1, data_index, dlc, lc, p, k, force)
            for column, value in zip(mark_data[mark], values):
                mark_data[mark][column].append(value)

        for i,data in enumerate(self):
            if f_index is not None and i > f_index:
                break
            dlc,lc,p,f,k=[],[],[],[],[]
            if data['artificial_judge']:
                # Export must not populate rawdata or reuse processed-data caches in the archive.
                fc = forcecurve()
                fc.data = copy.deepcopy(data)
                fc.recover_force(ljp)
                data_y = fc.get_prodata()['retract']['vDeflection']*1e12
                peaks = fc.data.get('peakindex', [])
                marks = fc.data.get('mark', [])
                dlc_values = fc.data.get('dlc', [])
                wlc_values = fc.data.get('wlcarg', [])
                k_values = fc.data.get('k', [])
                highspeed = fc.data.get('offset', {}).get('highspeed', 0) * 1e12
                mark_filter = fc.data.get('arg', {}).get('mark', {})
                force_filter = {}
                for mk, rng in mark_filter.items():
                    n = len(rng)
                    if (n >= 3 and rng[2] is not None) or (n >= 4 and rng[3] is not None):
                        force_filter[mk] = rng
                for index,p_i in enumerate(peaks):
                    force_value = data_y[:,0][p_i] + highspeed
                    dlc_value = dlc_values[index] if index < len(dlc_values) else None
                    peak_mark = marks[index] if index < len(marks) else ('none' if dlc_value is not None else None)
                    wlc_value = wlc_values[index] if index < len(wlc_values) else ()
                    lc_value = wlc_value[0] if len(wlc_value) > 0 else None
                    p_value = wlc_value[1] if len(wlc_value) > 1 else None
                    k_value = k_values[index] if index < len(k_values) else None
                    if peak_mark in force_filter:
                        rang = force_filter[peak_mark]
                        min_f = rang[2] if len(rang) >= 3 else None
                        max_f = rang[3] if len(rang) >= 4 else None
                        skip = False
                        if min_f is not None and force_value < min_f:
                            skip = True
                        if max_f is not None and force_value > max_f:
                            skip = True
                        if skip:
                            peak_mark = 'none'
                    if peak_mark is not None:
                        append_mark(peak_mark, i, index, p_i, dlc_value,
                                    lc_value, p_value, k_value, force_value)
                    if dlc_value is not None:
                        dlc.append(dlc_value)
                    lc.append(lc_value)
                    p.append(p_value)
                    k.append(k_value)
                    f.append(force_value)
                if include_last and peaks:
                    index = len(peaks) - 1
                    p_i = peaks[index]
                    wlc_value = wlc_values[index] if index < len(wlc_values) else ()
                    lc_value = wlc_value[0] if len(wlc_value) > 0 else None
                    p_value = wlc_value[1] if len(wlc_value) > 1 else None
                    k_value = k_values[index] if index < len(k_values) else None
                    force_value = data_y[:,0][p_i] + highspeed
                    append_mark('last_peak', i, index, p_i, None,
                                lc_value, p_value, k_value, force_value)
            arg_dic['dlc'].append(dlc)
            arg_dic['k'].append(k)
            arg_dic['force'].append(f)
            arg_dic['lc'].append(lc)
            arg_dic['p'].append(p)
            arg_dic['lens'].append(len(f))
            if f_index!=None and i==f_index:
                break
        max_len = max(arg_dic['lens'])
        for i,lens in enumerate(arg_dic['lens']):
            n = max_len - lens
            arg_dic['dlc'][i]+=[None]*n
            arg_dic['k'][i]+=[None]*n
            arg_dic['force'][i]+=[None]*n
            arg_dic['lc'][i]+=[None]*n
            arg_dic['p'][i]+=[None]*n
        del arg_dic['lens']

        try:
            from openpyxl import Workbook
            wb = Workbook()
            wb.remove(wb.active)
            for sheet_name, data in arg_dic.items():
                ws = wb.create_sheet(title=sheet_name[:31])
                ncols = max((len(r) for r in data), default=0)
                for ri in range(ncols):
                    ws.cell(row=1, column=ri+1, value=ri)
                for ci, row in enumerate(data):
                    for ri, val in enumerate(row):
                        cell = ws.cell(row=ci+2, column=ri+1)
                        cell.value = round(val, 2) if isinstance(val, float) else val
            wb.save(os.path.join(os.path.dirname(self.fname), "INDEX-{}.xlsx".format(os.path.splitext(os.path.basename(self.fname))[0])))
        except:
            import traceback
            traceback.print_exc()
            return False

        try:
            from openpyxl import Workbook
            wb = Workbook()
            wb.remove(wb.active)
            for mark, item in mark_data.items():
                ws = wb.create_sheet(title=mark[:31])
                cols = list(item.keys())
                for ci, col_name in enumerate(cols):
                    ws.cell(row=1, column=ci+1, value=col_name)
                    for ri, val in enumerate(item[col_name]):
                        cell = ws.cell(row=ri+2, column=ci+1)
                        cell.value = round(val, 2) if isinstance(val, float) else val
            wb.save(os.path.join(os.path.dirname(self.fname), "MARK-{}.xlsx".format(os.path.splitext(os.path.basename(self.fname))[0])))
        except:
            import traceback
            traceback.print_exc()
            return False
        return True
            
    def extrac_argdata(self, ljp, filters=True, filter_lst=['peaknum_judge', 'mobilenet_judge', 'artificial_judge']):
        wk_i = xlwt.Workbook(encoding='utf-8')
        ws_i_lc = wk_i.add_sheet('lc')
        ws_i_dlc = wk_i.add_sheet('dlc')
        ws_i_lp = wk_i.add_sheet('lp')
        ws_i_force = wk_i.add_sheet('force')
        wk_m = xlwt.Workbook(encoding='utf-8')
        ws_m_lc = wk_m.add_sheet('lc')
        ws_m_dlc = wk_m.add_sheet('dlc')
        ws_m_lp = wk_m.add_sheet('lp')
        ws_m_force = wk_m.add_sheet('force')
        col_count = 0
        mk_lst = []
        mk_conut = {}
        for i in range(len(self)):
            fc1 = forcecurve()
            fc1.data = self[i]
            fc1.recover_force(ljp)
            if filters and sum([fc1.data[t] for t in filter_lst]) < len(filter_lst):
                continue
            force = fc1.get_prodata()['retract']['vDeflection'].reshape(-1)[fc1.data['peakindex']] * 1e12
            lc, lp = np.array(fc1.data['wlcarg'])[:, 0], np.array(fc1.data['wlcarg'])[:, 1]
            dlc = fc1.data['dlc']
            mark = fc1.data['mark']
            for z in range(len(force)):
                ws_i_lc.write(col_count, z, lc[z])
                if z < len(force) - 1:
                    ws_i_dlc.write(col_count, z, dlc[z])
                ws_i_lp.write(col_count, z, lp[z])
                ws_i_force.write(col_count, z, force[z])
            for z in range(len(mark)):
                if mark[z] not in mk_lst:
                    mk_lst.append(mark[z])
                    mk_conut[mark[z]] = 1
                    for ws in [ws_m_lc, ws_m_dlc, ws_m_lp, ws_m_force]:
                        ws.write(0, mk_lst.index(mark[z]), mark[z])
                else:
                    mk_conut[mark[z]] += 1
                ws_m_lc.write(mk_conut[mark[z]], mk_lst.index(mark[z]), lc[z])
                ws_m_dlc.write(mk_conut[mark[z]], mk_lst.index(mark[z]), dlc[z])
                ws_m_lp.write(mk_conut[mark[z]], mk_lst.index(mark[z]), lp[z])
                ws_m_force.write(mk_conut[mark[z]], mk_lst.index(mark[z]), force[z])
            col_count += 1
        todir = os.path.dirname(self.fname)
        wk_i_svname = os.path.join(todir, 'outputdata base on index of {}.xls'.format(os.path.splitext(os.path.basename(self.fname))[0]))
        wk_m_svname = os.path.join(todir, 'outputdata base on mark of {}.xls'.format(os.path.splitext(os.path.basename(self.fname))[0]))
        try:
            wk_i.save(wk_i_svname)
            wk_m.save(wk_m_svname)
        except:
            return False
        return True
    def export_celldata(self,ljp,f_index):
        fc = forcecurve()
        dic = {'abs force':[],'force':[],'k':[],'lens':[]}
        for i,data in enumerate(self):
            fc.data = data
            fc.recover_force(ljp)
            data = fc.get_prodata()['retract']
            data_x,data_y = data['measuredHeight']*1e9,data['vDeflection']*1e12
            af,f,k=[],[],[]
            if not fc.data['artificial_judge']:
                dic['abs force'].append(af)
                dic['force'].append(f)
                dic['k'].append(k)
                dic['lens'].append(0)
                continue
            for i,p_i in  enumerate(fc.data['peakindex']):
                af.append(data_y[p_i][0]+fc.data['offset']['highspeed']*1e12)
                b_i = fc.data['bottomindex'][np.argmin(np.abs(p_i-fc.data['bottomindex']))]
                f.append((data_y[p_i]-data_y[b_i])[0]+fc.data['offset']['highspeed']*1e12)
                k.append(fc.data['k'][i])
            dic['abs force'].append(af)
            dic['force'].append(f)
            dic['k'].append(k)
            dic['lens'].append(len(fc.data['peakindex']))
            if i==f_index:
                break
        max_lens = max(dic['lens'])
        for i,lens in enumerate(dic['lens']):
            dic['abs force'][i] += [None] * (max_lens - lens)
            dic['force'][i] += [None] * (max_lens - lens)
            dic['k'][i] += [None] * (max_lens - lens)
        try:
            from openpyxl import Workbook
            wb = Workbook()
            wb.remove(wb.active)
            for sheet_name, data in dic.items():
                if sheet_name == 'lens':
                    continue
                ws = wb.create_sheet(title=sheet_name[:31])
                ncols = max((len(r) for r in data), default=0)
                for ri in range(ncols):
                    ws.cell(row=1, column=ri+1, value=ri)
                for ci, row in enumerate(data):
                    for ri, val in enumerate(row):
                        cell = ws.cell(row=ci+2, column=ri+1)
                        cell.value = round(val, 2) if isinstance(val, float) else val
            wb.save(os.path.join(os.path.dirname(self.fname), "cell_curve-{}.xlsx".format(os.path.splitext(os.path.basename(self.fname))[0])))
        except:
            return False
        return True
